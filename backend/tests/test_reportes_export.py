import io
import uuid

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def _unique(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:8]}"


def _seed(client: TestClient, mobile_auth_headers: dict) -> dict:
    dep = client.post(
        "/api/v1/depositos",
        json={"nombre": _unique("DepExp")},
        headers=mobile_auth_headers,
    ).json()
    sec = client.post(
        f"/api/v1/depositos/{dep['id']}/sectores",
        json={"nombre": "Sector Exp"},
        headers=mobile_auth_headers,
    ).json()
    ubi = client.post(
        f"/api/v1/depositos/{dep['id']}/sectores/{sec['id']}/ubicaciones",
        json={"codigo": "E-01"},
        headers=mobile_auth_headers,
    ).json()
    cat = client.post(
        "/api/v1/categorias",
        json={"nombre": _unique("CatExp")},
        headers=mobile_auth_headers,
    ).json()
    epc = f"E280EXP{_unique('')[:6]}".upper()
    activo = client.post(
        "/api/v1/activos",
        json={
            "numero_patrimonial": _unique("PAT-E"),
            "descripcion": "Activo exportable",
            "categoria_id": cat["id"],
            "epc": epc,
        },
        headers=mobile_auth_headers,
    ).json()
    client.post(
        f"/api/v1/activos/{activo['id']}/asignar-ubicacion",
        json={"ubicacion_id": ubi["id"]},
        headers=mobile_auth_headers,
    )
    inv = client.post(
        "/api/v1/inventarios",
        json={"deposito_id": dep["id"]},
        headers=mobile_auth_headers,
    ).json()
    client.post(
        f"/api/v1/inventarios/{inv['id']}/cerrar",
        json={"epcs": [epc]},
        headers=mobile_auth_headers,
    )
    return {"deposito": dep, "activo": activo, "inventario": inv, "epc": epc}


def test_export_movimientos_xlsx(client: TestClient, mobile_auth_headers, auth_headers):
    _seed(client, mobile_auth_headers)
    response = client.get(
        "/api/v1/reportes/movimientos",
        params={"formato": "xlsx"},
        headers=auth_headers,
    )
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    assert "attachment" in response.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(response.content))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "accion" in headers
    assert "numero_patrimonial" in headers
    assert ws.max_row >= 2


def test_export_stock_csv(client: TestClient, mobile_auth_headers, auth_headers):
    seed = _seed(client, mobile_auth_headers)
    response = client.get(
        f"/api/v1/reportes/stock/{seed['deposito']['id']}",
        params={"formato": "csv"},
        headers=auth_headers,
    )
    assert response.status_code == 200
    assert "text/csv" in response.headers["content-type"]
    text = response.content.decode("utf-8-sig")
    assert "numero_patrimonial" in text
    assert seed["activo"]["numero_patrimonial"] in text


def test_export_inventario_pdf_y_xlsx(client: TestClient, mobile_auth_headers, auth_headers):
    seed = _seed(client, mobile_auth_headers)
    inv_id = seed["inventario"]["id"]

    xlsx = client.get(
        f"/api/v1/reportes/inventarios/{inv_id}",
        params={"formato": "xlsx"},
        headers=auth_headers,
    )
    assert xlsx.status_code == 200, xlsx.text
    wb = load_workbook(io.BytesIO(xlsx.content))
    assert "Faltantes" in wb.sheetnames
    assert "Encontrados" in wb.sheetnames

    pdf = client.get(
        f"/api/v1/reportes/inventarios/{inv_id}",
        params={"formato": "pdf"},
        headers=auth_headers,
    )
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content[:4] == b"%PDF"


def test_export_formato_invalido(client: TestClient, mobile_auth_headers, auth_headers):
    response = client.get(
        "/api/v1/reportes/movimientos",
        params={"formato": "docx"},
        headers=auth_headers,
    )
    assert response.status_code == 400
